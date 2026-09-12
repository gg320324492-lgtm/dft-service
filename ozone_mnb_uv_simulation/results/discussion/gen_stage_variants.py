# 第3页 四阶段任务页：三种布局方案（A 横向四卡 / B 双列卡片 / C 垂直时间线）
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TPL = r'C:\Users\pc\WPSDrive\757644119\WPS云盘\综合文件夹\微纳米气泡\汇报\2026.6.18汇报\2026-6-18-杜同贺.pptx'
OUT = r'results/discussion/assets/render/stage_variants.pptx'

NAVY = '002060'; RED = 'C00000'; BLACK = '000000'
BAND_BLUE = '5B9BD5'; LIGHT_BLUE = 'DCE9F5'; TINT2 = 'F5FAFB'
HAIR = 'C9DAE0'; MUTED = '5B7080'; PRIMARY = '002060'
LIGHT_TXT = 'E8F1F8'; CREAM = 'FFE3D1'; ACCENT = 'C00000'
F = '微软雅黑'
TITLE = dict(size=20, bold=True, color=NAVY)
BODY = dict(size=12, bold=True, color=BLACK)

prs = Presentation(TPL)
n_orig = len(prs.slides)

def clone_slide(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    idmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        if rel.is_external:
            idmap[rId] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            idmap[rId] = dst.part.relate_to(rel.target_part, rel.reltype)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for e in el.iter():
            for a in ('embed', 'link', 'id'):
                k = qn('r:' + a)
                v = e.get(k)
                if v and v in idmap:
                    e.set(k, idmap[v])
        dst.shapes._spTree.append(el)
    return dst

def rm_shapes(slide, keep_names):
    for sh in list(slide.shapes):
        if sh.name not in keep_names:
            sh._element.getparent().remove(sh._element)

def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh

def style_run(r, size=None, bold=True, color=BLACK, latin='Times New Roman', ea=F, baseline=None):
    if size is not None:
        r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = latin
    rPr = r._r.get_or_add_rPr()
    if baseline is not None:
        rPr.set('baseline', str(baseline))
    latin_el = rPr.find(qn('a:latin'))
    ea_el = rPr.find(qn('a:ea'))
    if ea_el is None:
        ea_el = rPr.makeelement(qn('a:ea'), {})
        if latin_el is not None:
            latin_el.addnext(ea_el)
        else:
            rPr.append(ea_el)
    ea_el.set('typeface', ea)

import re as _re
def fx_parse(s):
    pat = _re.compile(r'([_^])\{([^}]*)\}')
    out, last = [], 0
    for m in pat.finditer(s):
        if m.start() > last:
            out.append((s[last:m.start()], None))
        out.append((m.group(2), m.group(1)))
        last = m.end()
    if last < len(s):
        out.append((s[last:], None))
    return out

def add_runs(para, text, **kw):
    for t, kind in fx_parse(text):
        r = para.add_run()
        r.text = t
        style_run(r, baseline=(-25000 if kind == '_' else 30000 if kind == '^' else None), **kw)

def add_textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        para.space_after = Pt(p.get('space_after', 3))
        for t, kw in p['runs']:
            add_runs(para, t, **kw)
    return tb

def retext(shape, paras):
    tf = shape.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', para.alignment)
        para.space_after = Pt(p.get('space_after', 3))
        for t, kw in p['runs']:
            add_runs(para, t, **kw)

def pill(slide, x, y, w, h, text, fill, tcolor):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    add_runs(para, text, size=10, bold=True, color=tcolor)

tpl_text = prs.slides[4]

STAGES = [
    ('① 臭氧水合基线', '已建立', 'done',
     '固定计算方法与验收标准，算准 H₂O、O₃、O₂ 在气相和 SMD 水相的结构、能量与频率。',
     '为后续所有反应计算提供参照点与统一方法口径（水相仍有未验收项）。'),
    ('② 目标物反应模型', '当前阶段', 'cur',
     '研究氨氮与 O₃、活性氧的复合物/中间体——当前是 NH₃ 与 O₃ 的稳定接触构型。',
     '搞清反应起点：两分子如何结合、结合多强（已获 2 个驻点候选，核查中）。'),
    ('③ 反应路径与能垒', '尚未开展', 'todo',
     '对优先反应路径做过渡态（TS）、频率与 IRC 验证，算出反应能和能垒。',
     '回答反应走哪条路、哪一步最慢、需要多大能量——直接对应工艺条件选择。'),
    ('④ UV 激发态·气液界面', '尚未开展', 'todo',
     '研究 UV 激发态（TD-DFT/光化学）与气液界面（微纳米气泡表面）的分子模型。',
     '解释紫外增强效果与气泡界面提高利用率的原因，须结合实验条件论证后立项。'),
]
PILL = {'done': (BAND_BLUE, 'FFFFFF'), 'cur': (ACCENT, 'FFFFFF'), 'todo': (HAIR, MUTED)}

def base():
    s = clone_slide(tpl_text)
    rm_shapes(s, {'矩形 10', '图片 2', '直接连接符 1', 'Text Box 2'})
    retext(find(s, 'Text Box 2'), [{'runs': [('四阶段计划：每个阶段的任务是干什么的', TITLE)]}])
    return s

# ---- 方案 A：横向四卡（路线图式） ----
sA = base()
for i, (name, status, kind, task, goal) in enumerate(STAGES):
    x = 0.35 + i * 2.35
    cur = kind == 'cur'
    card = sA.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.15), Inches(2.25), Inches(4.35))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string(NAVY if cur else TINT2)
    card.line.color.rgb = RGBColor.from_string(NAVY if cur else HAIR)
    card.line.width = Pt(1.25 if cur else 0.75)
    card.shadow.inherit = False
    pill(sA, x + 0.55, 1.31, 1.15, 0.3, status, ACCENT if cur else BAND_BLUE, 'FFFFFF')
    add_textbox(sA, x + 0.1, 1.68, 2.05, 0.75,
                [{'runs': [(name, dict(size=13, bold=True, color='FFFFFF' if cur else PRIMARY))], 'align': PP_ALIGN.CENTER}])
    ln = sA.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.25), Inches(2.42), Inches(1.75), Pt(1))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGBColor.from_string('FFFFFF' if cur else HAIR)
    ln.line.fill.background()
    ln.shadow.inherit = False
    add_textbox(sA, x + 0.15, 2.56, 1.95, 2.85, [
        {'runs': [('任务：', dict(size=11, bold=True, color=ACCENT if cur else NAVY)),
                  (task, dict(size=11, bold=True, color=LIGHT_TXT if cur else BLACK))], 'space_after': 8},
        {'runs': [('目的：', dict(size=11, bold=True, color=ACCENT if cur else NAVY)),
                  (goal, dict(size=11, bold=True, color=LIGHT_TXT if cur else BLACK))]},
    ])

# ---- 方案 B：双列 2×2 卡片 ----
sB = base()
posB = [(0.35, 1.12), (5.1, 1.12), (0.35, 3.78), (5.1, 3.78)]
for i, (name, status, kind, task, goal) in enumerate(STAGES):
    x, y = posB[i]
    cur = kind == 'cur'
    card = sB.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(4.55), Inches(2.5))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string(NAVY if cur else TINT2)
    card.line.color.rgb = RGBColor.from_string(NAVY if cur else HAIR)
    card.line.width = Pt(1.25 if cur else 0.75)
    card.shadow.inherit = False
    add_textbox(sB, x + 0.2, y + 0.14, 3.0, 0.4,
                [{'runs': [(name, dict(size=13.5, bold=True, color='FFFFFF' if cur else PRIMARY))]}])
    pill(sB, x + 3.25, y + 0.16, 1.1, 0.3, status, ACCENT if cur else BAND_BLUE, 'FFFFFF')
    add_textbox(sB, x + 0.2, y + 0.66, 4.15, 1.7, [
        {'runs': [('任务：', dict(size=12, bold=True, color=CREAM if cur else NAVY)),
                  (task, dict(size=12, bold=True, color=LIGHT_TXT if cur else BLACK))], 'space_after': 6},
        {'runs': [('目的：', dict(size=12, bold=True, color=CREAM if cur else NAVY)),
                  (goal, dict(size=12, bold=True, color=LIGHT_TXT if cur else BLACK))]},
    ])

# ---- 方案 C：垂直时间线 ----
sC = base()
ln = sC.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.98), Inches(1.45), Pt(2.5), Inches(3.95))
ln.fill.solid()
ln.fill.fore_color.rgb = RGBColor.from_string(BAND_BLUE)
ln.line.fill.background()
ln.shadow.inherit = False
for i, (name, status, kind, task, goal) in enumerate(STAGES):
    y = 1.15 + i * 1.1
    cur = kind == 'cur'
    node = sC.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72), Inches(y + 0.02), Inches(0.52), Inches(0.52))
    node.fill.solid()
    node.fill.fore_color.rgb = RGBColor.from_string(NAVY if cur else BAND_BLUE)
    node.line.fill.background()
    node.shadow.inherit = False
    ntf = node.text_frame
    ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
    ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np = ntf.paragraphs[0]
    np.alignment = PP_ALIGN.CENTER
    add_runs(np, ['①', '②', '③', '④'][i], size=14, bold=True, color='FFFFFF')
    add_textbox(sC, 1.55, y, 8.2, 0.38, [
        {'runs': [(name, dict(size=13.5, bold=True, color=NAVY)),
                  ('　' + status, dict(size=11.5, bold=True, color=ACCENT if cur else MUTED))]},
    ])
    add_textbox(sC, 1.55, y + 0.42, 8.2, 0.68, [
        {'runs': [('任务：', dict(size=12, bold=True, color=NAVY)), (task, BODY)], 'space_after': 2},
        {'runs': [('目的：', dict(size=12, bold=True, color=NAVY)), (goal, BODY)]},
    ])

# 删除原始模板页
for i in range(n_orig - 1, -1, -1):
    sld = list(prs.slides._sldIdLst)[i]
    prs.part.drop_rel(sld.get(qn('r:id')))
    prs.slides._sldIdLst.remove(sld)

prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
