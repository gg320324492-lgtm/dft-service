# 第2页 目标链条注释：三种布局方案（A 芯片下注释 / B 双层卡片 / C 整条注释带）
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TPL = r'C:\Users\pc\WPSDrive\757644119\WPS云盘\综合文件夹\微纳米气泡\汇报\2026.6.18汇报\2026-6-18-杜同贺.pptx'
OUT = r'results/discussion/assets/render/slide2_variants.pptx'

NAVY = '002060'; RED = 'C00000'; BLACK = '000000'
BAND_BLUE = '5B9BD5'; LIGHT_BLUE = 'DCE9F5'; TINT2 = 'F5FAFB'
HAIR = 'C9DAE0'; MUTED = '5B7080'; PRIMARY = '002060'
F = '微软雅黑'
TITLE = dict(size=20, bold=True, color=NAVY)
BODY = dict(size=12, bold=True, color=BLACK)

prs = Presentation(TPL)

def clone_slide(src):
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    idmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith('/slideLayout') or rel.reltype.endswith('/notesSlide'):
            continue
        if rel.is_external:
            idmap[rId] = dst.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            idmap[rId] = dst.part.relate_to(rel.target_part, rel.reltype)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        for e in el.iter():
            for a in ('embed', 'link', 'id'):
                k = qn('r:' + a)
                v = e.get(k)
                if v and v in idmap:
                    e.set(k, idmap[v])
        dst.shapes._spTree.append(el)
    return dst

def rm_shapes(slide, keep_names):
    for sh in list(slide.shapes):
        if sh.name not in keep_names:
            sh._element.getparent().remove(sh._element)

def find(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh

def style_run(r, size=None, bold=True, color=BLACK, latin='Times New Roman', ea=F, baseline=None):
    if size is not None:
        r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    r.font.name = latin
    rPr = r._r.get_or_add_rPr()
    if baseline is not None:
        rPr.set('baseline', str(baseline))
    latin_el = rPr.find(qn('a:latin'))
    ea_el = rPr.find(qn('a:ea'))
    if ea_el is None:
        ea_el = rPr.makeelement(qn('a:ea'), {})
        if latin_el is not None:
            latin_el.addnext(ea_el)
        else:
            rPr.append(ea_el)
    ea_el.set('typeface', ea)

import re as _re
def fx_parse(s):
    pat = _re.compile(r'([_^])\{([^}]*)\}')
    out, last = [], 0
    for m in pat.finditer(s):
        if m.start() > last:
            out.append((s[last:m.start()], None))
        out.append((m.group(2), m.group(1)))
        last = m.end()
    if last < len(s):
        out.append((s[last:], None))
    return out

def add_runs(para, text, **kw):
    for t, kind in fx_parse(text):
        r = para.add_run()
        r.text = t
        style_run(r, baseline=(-25000 if kind == '_' else 30000 if kind == '^' else None), **kw)

def add_textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        para.space_after = Pt(p.get('space_after', 4))
        for t, kw in p['runs']:
            add_runs(para, t, **kw)
    return tb

def retext(shape, paras):
    tf = shape.text_frame
    for p in list(tf.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    for r in list(tf.paragraphs[0].runs):
        r._r.getparent().remove(r._r)
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', para.alignment)
        para.space_after = Pt(p.get('space_after', 4))
        for t, kw in p['runs']:
            add_runs(para, t, **kw)

def chip(slide, x, y, w, h, text, fill, tcolor, tsize=12.5):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    c.fill.solid()
    c.fill.fore_color.rgb = RGBColor.from_string(fill)
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    add_runs(para, text, size=tsize, bold=True, color=tcolor)
    return c

tpl_text = prs.slides[4]

# 原始模板页稍后统一删除
n_orig = len(prs.slides)

bg_lines = [
    '背景：养殖水体氨氮超标，对鱼有毒性，也是水产养殖水质管理的核心指标之一。',
    '工艺思路：以臭氧微纳米气泡与紫外联用的方式去除氨氮。',
    '关键问题：它为什么有效？臭氧和氨氮到底是怎么反应的？',
    '实验的局限：只能测得去除率，难以直接观测分子层面的反应过程。',
]
cap = ['臭氧溶入水体\n紫外分解产生活性氧',
       '·OH 等强氧化性物种\n氧化水中氨氮',
       '氨氮被活性氧\n逐步氧化为中间产物',
       '最终氧化产物\n可由水质检测直接监测']
names = [('O₃ / UV 活化', 1.85), ('活性氧 ROS', 1.6), ('NH₃ / NH₄⁺ 氧化', 2.05), ('NH₂OH · NO₂⁻ · NO₃⁻', 2.9)]

def base(variant_name):
    s = clone_slide(tpl_text)
    rm_shapes(s, {'矩形 10', '图片 2', '直接连接符 1', 'Text Box 2'})
    retext(find(s, 'Text Box 2'), [{'runs': [('研究背景与总体路线：为什么做分子模拟', TITLE)]}])
    add_textbox(s, 0.35, 1.06, 9.3, 1.45,
                [{'runs': [('• ' + t, BODY)], 'space_after': 6} for t in bg_lines])
    add_textbox(s, 0.35, 3.02, 4, 0.28, [{'runs': [('研究目标链条', dict(size=13, bold=True, color=NAVY))]}])
    add_textbox(s, 0.35, 4.86, 9.3, 0.6,
                [{'runs': [('实验测得处理效果，分子模拟解释背后机理——说明臭氧与氨氮在分子层面如何反应，为工艺优化提供依据。',
                            dict(size=12, bold=True, color=NAVY))]}])
    return s

def add_arrows(s, xs, ws, y):
    cx = 0.35
    for i in range(3):
        cx += ws[i]
        add_textbox(s, cx - 0.02, y + 0.04, 0.28, 0.35,
                    [{'runs': [('→', dict(size=13, bold=True, color=NAVY))], 'align': PP_ALIGN.CENTER}])
        cx += 0.28

# ---- 方案 A：芯片 + 下方小字注释 ----
sA = base('A')
cy = 3.4
cx = 0.35
for i, (t, w) in enumerate(names):
    chip(sA, cx, cy, w, 0.5, t, BAND_BLUE, 'FFFFFF', 12.5)
    add_textbox(sA, cx - 0.08, cy + 0.58, w + 0.16, 0.6,
                [{'runs': [(cap[i].split('\n')[0], dict(size=10.5, color=MUTED))], 'align': PP_ALIGN.CENTER, 'space_after': 0},
                 {'runs': [(cap[i].split('\n')[1], dict(size=10.5, color=MUTED))], 'align': PP_ALIGN.CENTER}])
    cx += w + 0.28
# 箭头
acc = 0.35
for i, (t, w) in enumerate(names):
    acc += w
    if i < 3:
        add_textbox(sA, acc - 0.02, 3.44, 0.28, 0.35,
                    [{'runs': [('→', dict(size=13, bold=True, color=NAVY))], 'align': PP_ALIGN.CENTER}])
        acc += 0.28

# ---- 方案 B：双层卡片（标题+注释在卡内） ----
sB = clone_slide(tpl_text)
rm_shapes(sB, {'矩形 10', '图片 2', '直接连接符 1', 'Text Box 2'})
retext(find(sB, 'Text Box 2'), [{'runs': [('研究背景与总体路线：为什么做分子模拟', TITLE)]}])
add_textbox(sB, 0.35, 1.06, 9.3, 1.45,
            [{'runs': [('• ' + t, BODY)], 'space_after': 6} for t in bg_lines])
add_textbox(sB, 0.35, 3.02, 4, 0.28, [{'runs': [('研究目标链条', dict(size=13, bold=True, color=NAVY))]}])
add_textbox(sB, 0.35, 4.86, 9.3, 0.6,
            [{'runs': [('实验测得处理效果，分子模拟解释背后机理——说明臭氧与氨氮在分子层面如何反应，为工艺优化提供依据。',
                        dict(size=12, bold=True, color=NAVY))]}])
cx = 0.35
for i, (t, w) in enumerate(names):
    card = sB.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(3.4), Inches(w), Inches(1.25))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor.from_string(BAND_BLUE)
    card.line.fill.background()
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    add_runs(p1, t, size=12.5, bold=True, color='FFFFFF')
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(3)
    for j, seg in enumerate(cap[i].split('\n')):
        pp = p2 if j == 0 else tf.add_paragraph()
        if j > 0:
            pp.alignment = PP_ALIGN.CENTER
        add_runs(pp, seg, size=10.5, color='E8F1F8')
    cx += w + 0.28

# ---- 方案 C：芯片行 + 整条注释带 ----
sC = clone_slide(tpl_text)
rm_shapes(sC, {'矩形 10', '图片 2', '直接连接符 1', 'Text Box 2'})
retext(find(sC, 'Text Box 2'), [{'runs': [('研究背景与总体路线：为什么做分子模拟', TITLE)]}])
add_textbox(sC, 0.35, 1.06, 9.3, 1.45,
            [{'runs': [('• ' + t, BODY)], 'space_after': 6} for t in bg_lines])
add_textbox(sC, 0.35, 2.98, 4, 0.28, [{'runs': [('研究目标链条', dict(size=13, bold=True, color=NAVY))]}])
cx = 0.35
for i, (t, w) in enumerate(names):
    chip(sC, cx, 3.44, w, 0.5, t, BAND_BLUE, 'FFFFFF', 12.5)
    cx += w + 0.28
strip = sC.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.35), Inches(4.08), Inches(9.3), Inches(0.58))
strip.fill.solid()
strip.fill.fore_color.rgb = RGBColor.from_string(LIGHT_BLUE)
strip.line.fill.background()
strip.shadow.inherit = False
tfs = strip.text_frame
tfs.word_wrap = True
tfs.margin_left = tfs.margin_right = Inches(0.12)
tfs.margin_top = tfs.margin_bottom = Inches(0.03)
tfs.vertical_anchor = MSO_ANCHOR.MIDDLE
paraS = tfs.paragraphs[0]
paraS.alignment = PP_ALIGN.CENTER
add_runs(paraS, '臭氧溶入水体并经紫外分解产生活性氧 → ·OH 等强氧化性物种将氨氮逐步氧化 → 依次生成 NH₂OH、NO₂⁻、NO₃⁻，可由水质检测直接监测',
         size=11.5, bold=True, color=NAVY)
add_textbox(sC, 0.35, 4.86, 9.3, 0.6,
            [{'runs': [('实验测得处理效果，分子模拟解释背后机理——说明臭氧与氨氮在分子层面如何反应，为工艺优化提供依据。',
                        dict(size=12, bold=True, color=NAVY))]}])

for i in range(n_orig - 1, -1, -1):
    sld = list(prs.slides._sldIdLst)[i]
    prs.part.drop_rel(sld.get(qn('r:id')))
    prs.slides._sldIdLst.remove(sld)

prs.save(OUT)
print('saved:', OUT, 'slides:', len(prs.slides))
